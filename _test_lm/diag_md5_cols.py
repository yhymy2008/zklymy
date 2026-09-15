"""诊断 out_Oneid_md5.xlsx：查找 md5 相关列，统计目标 md5 出现位置。"""
from openpyxl import load_workbook

PATH = r"c:\Users\CNYangMe8\Documents\work\AI_coding\out_Oneid_md5.xlsx"
TARGET = "0000b11b600609f69c45178133be1829"

wb = load_workbook(PATH, read_only=True, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    hdr = next(ws.iter_rows(values_only=True))
    md5_cols = [i for i, c in enumerate(hdr) if c is not None and "md5" in str(c).lower()]
    print(f"== sheet: {name} | 列数: {len(hdr)} | md5相关列: {[(i, hdr[i]) for i in md5_cols]}")
    for i in md5_cols:
        cnt = 0
        rows = []
        for rno, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            v = r[i] if i < len(r) else None
            if v is not None and str(v).strip().lower() == TARGET:
                cnt += 1
                rows.append(rno)
        print(f"   列[{i}]({hdr[i]}) 含目标md5行数: {cnt} 行号: {rows[:10]}")
    # 也统计整表所有列中目标 md5 的出现（找非 md5 命名列）
    if not md5_cols:
        found = []
        for rno, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            for j, v in enumerate(r):
                if v is not None and str(v).strip().lower() == TARGET:
                    found.append((rno, j))
        print(f"   全表搜索目标md5: {len(found)} 处 {found[:10]}")
wb.close()
