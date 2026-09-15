"""快速搜索：只读最近修改的 xlsx/csv，遇到目标 md5 立即输出。"""
import os
import sys
import time

target = "0000b11b600609f69c45178133be1829"
search_roots = [
    r"C:\Users\CNYangMe8\Desktop",
    r"C:\Users\CNYangMe8\Downloads",
    r"C:\Users\CNYangMe8\Documents\work\AI_coding",
]
excl_dirs = {"node_modules", ".venv", ".git", "AppData", "Windows", "Program Files",
             "Program Files (x86)", "support-mini-program", "__pycache__"}
cutoff = time.time() - 30 * 86400  # 近 30 天

from openpyxl import load_workbook

hits = 0
checked = 0
for root in search_roots:
    if not os.path.isdir(root):
        continue
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if d not in excl_dirs and not d.startswith(".")]
        for f in fns:
            low = f.lower()
            if not low.endswith((".xlsx", ".xls", ".csv")):
                continue
            p = os.path.join(dp, f)
            try:
                if os.path.getmtime(p) < cutoff:
                    continue
            except OSError:
                continue
            checked += 1
            try:
                if low.endswith(".csv"):
                    with open(p, "r", encoding="utf-8-sig", errors="replace") as fh:
                        for line in fh:
                            if target in line:
                                print("[csv]", p, "->", line.strip()[:120], flush=True)
                                hits += 1
                                break
                else:
                    wb = load_workbook(p, read_only=True, data_only=True)
                    for name in wb.sheetnames:
                        ws = wb[name]
                        for r in ws.iter_rows(values_only=True):
                            for v in r:
                                if v is not None and str(v).strip().lower() == target:
                                    print("[xlsx]", p, "sheet=", name, flush=True)
                                    hits += 1
                                    break
                    wb.close()
            except Exception:
                pass
print(f"--- checked {checked} files, hits {hits}", flush=True)
