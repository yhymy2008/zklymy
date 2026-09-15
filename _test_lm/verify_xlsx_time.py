# -*- coding: utf-8 -*-
"""验证 xlsx 格式 + 显式 --a-time/--b-time 参数 + 时间差列"""
import os
import subprocess
import sys

from openpyxl import Workbook, load_workbook

TMP = r"c:\Users\CNYangMe8\Documents\work\AI_coding\_test_lm"
SCRIPT = r"c:\Users\CNYangMe8\Documents\work\AI_coding\lookup_merge.py"

# A: xlsx 含 Buytime
a_xlsx = os.path.join(TMP, "a_time.xlsx")
wb = Workbook(); ws = wb.active
ws.append(["手机号", "姓名", "Buytime"])
ws.append(["13800000001", "张三", "2026-01-01 10:00:00"])
ws.append(["13800000002", "李四", "2026-02-01 08:30:00"])
wb.save(a_xlsx)

# B: xlsx 含 event_time，13800000002 两条
b_xlsx = os.path.join(TMP, "b_time.xlsx")
wb = Workbook(); ws = wb.active
ws.append(["手机号", "城市", "event_time"])
ws.append(["13800000001", "上海", "2026-01-01 14:30:00"])
ws.append(["13800000002", "北京", "2026-02-01 10:00:00"])
ws.append(["13800000002", "北京", "2026-02-02 09:00:00"])
wb.save(b_xlsx)

out = os.path.join(TMP, "out_time_xlsx.xlsx")
cmd = [sys.executable, SCRIPT,
       "-a", a_xlsx, "-b", b_xlsx,
       "-ak", "手机号", "-bk", "手机号",
       "--a-time", "Buytime", "--b-time", "event_time",
       "-o", out]
env = dict(os.environ, PYTHONIOENCODING="utf-8")
r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", env=env)
print(r.stdout)
if r.returncode != 0:
    print("STDERR:", r.stderr)
    sys.exit(1)

wb = load_workbook(out, data_only=True)
ws = wb.active
data = [[c for c in row] for row in ws.iter_rows(values_only=True)]
for row in data:
    print("  ", row)

assert data[0] == ["手机号", "姓名", "Buytime", "城市", "event_time", "时间差_小时"], f"表头错误: {data[0]}"
assert len(data) == 4, f"行数错误: {len(data)}"  # 1 + 2 条李四 = 3 数据行 + 表头
assert data[1][5] == 4.5, f"张三差值: {data[1][5]}"
assert data[2][5] == 1.5 and data[3][5] == 24.5, f"李四差值: {data[2][5]}, {data[3][5]}"
print(">>> xlsx + 显式时间列参数断言全部通过")
