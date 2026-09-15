"""查找含 md5_hash 表头或目标 md5 值的文件（限定目录，快速）。"""
import os
from openpyxl import load_workbook

TARGET = "0000b11b600609f69c45178133be1829"
roots = [
    r"C:\Users\CNYangMe8\Documents\work\AI_coding\support-mini-program",
    r"C:\Users\CNYangMe8\Desktop",
]
excl = {"node_modules", ".venv", ".git", "__pycache__", "AppData"}

for root in roots:
    if not os.path.isdir(root):
        continue
    for dp, dns, fns in os.walk(root):
        dns[:] = [d for d in dns if d not in excl and not d.startswith(".")]
        for f in fns:
            low = f.lower()
            if not low.endswith((".xlsx", ".xls", ".csv")):
                continue
            p = os.path.join(dp, f)
            try:
                if low.endswith(".csv"):
                    with open(p, "r", encoding="utf-8-sig", errors="replace") as fh:
                        first = fh.readline().lower()
                        has_hdr = "md5_hash" in first or "md5" in first
                        has_t = False
                        for line in fh:
                            if TARGET in line:
                                has_t = True
                                break
                    if has_hdr or has_t:
                        print(f"[csv] {p} hdr_md5={has_hdr} target={has_t}")
                else:
                    wb = load_workbook(p, read_only=True, data_only=True)
                    for name in wb.sheetnames:
                        ws = wb[name]
                        hdr = next(ws.iter_rows(values_only=True))
                        hdr_low = [str(c).lower() for c in hdr if c is not None]
                        has_hdr = any("md5" in h for h in hdr_low)
                        has_t = False
                        if has_hdr:
                            md5_idxs = [i for i, h in enumerate(hdr_low) if "md5" in h]
                            for r in ws.iter_rows(min_row=2, values_only=True):
                                for i in md5_idxs:
                                    v = r[i] if i < len(r) else None
                                    if v is not None and str(v).strip().lower() == TARGET:
                                        has_t = True
                                        break
                                if has_t:
                                    break
                        if has_hdr or has_t:
                            print(f"[xlsx] {p} sheet={name} md5_cols={[hdr[i] for i in md5_idxs] if has_hdr else []} target={has_t}")
                    wb.close()
            except Exception:
                pass
print("done")
